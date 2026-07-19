from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TEMPLATE = Path("/Users/omarabdelaziz/Downloads/CF-paper-template (3).docx")
OUT = Path(
    "/Volumes/Extreme Pro/done_shop_projejcts/apps/glory_gym/"
    "Beta_Function_BAS115_Research_English.docx"
)


def clear_document(document: Document) -> None:
    body = document._body._element
    sect_pr = None
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            sect_pr = child
        body.remove(child)
    if sect_pr is not None:
        body.append(sect_pr)


def set_run_font(run, size=12.5, bold=False, italic=False, color=None) -> None:
    font = "Times New Roman"
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:cs"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def para(
    document,
    text="",
    *,
    size=12.5,
    bold=False,
    italic=False,
    align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    after=7,
    before=0,
    color=None,
):
    p = document.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, italic=italic, color=color)
    return p


def heading(document, text, *, level=1):
    size = 15 if level == 1 else 13.5
    return para(
        document,
        text,
        size=size,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        after=8,
        before=3,
    )


def bullet(document, text):
    p = para(document, "• " + text, size=12.2, after=4, align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.left_indent = Inches(0.22)
    return p


def page_break(document):
    p = document.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def border_cell(cell, color="B8B8B8"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "8")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def formula_table(document, rows):
    table = document.add_table(rows=len(rows), cols=2)
    table.autofit = False
    widths = (Inches(2.25), Inches(4.95))
    for row_index, row in enumerate(table.rows):
        for cell_index, cell in enumerate(row.cells):
            cell.width = widths[cell_index]
            border_cell(cell)
            if row_index == 0:
                shade_cell(cell, "EDEDED")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(rows[row_index][cell_index])
            set_run_font(run, size=11.5, bold=(row_index == 0))
    para(document, "", after=2)


def setup(document: Document):
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.68)
        section.bottom_margin = Inches(0.68)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)
        section.header_distance = Inches(0.32)
        section.footer_distance = Inches(0.32)

        header_p = section.header.paragraphs[0]
        header_p.clear()
        header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = header_p.add_run(
            "BAS115 Complex Functions Student Activity, Volume 2025-2026, Issue Spring, May 2026"
        )
        set_run_font(r, size=8.5, italic=True, color="555555")

        footer_p = section.footer.paragraphs[0]
        footer_p.clear()
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = footer_p.add_run("Beta Function Research")
        set_run_font(r, size=8.5, color="555555")

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12.5)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")


def build():
    document = Document(str(TEMPLATE))
    clear_document(document)
    setup(document)

    para(
        document,
        "Beta Function in Complex Functions",
        size=21,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=8,
    )
    para(
        document,
        "A paper submitted to BAS115 Complex Functions Student Activity",
        size=14,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=6,
    )
    para(
        document,
        "Student Name | Department, Institute Name",
        size=12,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=16,
    )

    para(document, "Abstract-", size=12.5, bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
    para(
        document,
        "This paper studies the Beta function B(m,n) as one of the most useful special functions in the course BAS115 Complex Functions. The discussion follows Lecture 6 and explains how the Beta function can be used to evaluate definite integrals in a shorter and more systematic way. The paper presents the three main forms of the Beta function, the relation between Beta and Gamma functions, and several solved applications. The main idea is that many integrals become easier when their powers and limits are matched with a standard Beta form.",
    )
    para(document, "Index Terms-", size=12.5, bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.LEFT, after=2)
    para(document, "Beta Function, Gamma Function, Special Functions, Definite Integrals, Complex Functions.", after=10)

    heading(document, "Introduction")
    para(
        document,
        "Special functions appear in mathematics when ordinary elementary methods are not enough to express or calculate an integral directly. The Beta function is a special function of two variables, usually written as B(m,n). It is important because it changes several difficult-looking integrals into one standard expression. Instead of finding an antiderivative every time, the student can compare the integral with a known Beta form and then use the properties of Beta and Gamma functions.",
    )
    para(
        document,
        "In BAS115 Complex Functions, the Beta function is introduced with the Gamma function because both of them are connected. This connection is very useful for integrals that contain powers, expressions such as 1-t or 1+t, and trigonometric factors involving sine and cosine. The purpose of this paper is to explain the topic clearly, starting from the definition and ending with applications similar to the lecture examples.",
    )
    para(
        document,
        "The importance of this topic also appears in the way it organizes problem solving. A normal integral may look complicated at first, but if it has the correct structure, it can be rewritten as a Beta function. This gives the student a clear path: identify the pattern, compare it with the formula, and then calculate the value. For this reason, the Beta function is not only a formula to memorize, but also a method of thinking about definite integrals.",
    )
    para(
        document,
        "This research follows the same sequence used in the lecture. First, it introduces the three standard forms of the Beta function. Second, it explains the problem and the method used to solve examples. Third, it studies the relation between Beta and Gamma functions. Finally, it gives numerical applications and a short conclusion about the value of the method.",
    )
    page_break(document)

    heading(document, "Definition and Standard Forms")
    para(
        document,
        "The first form of the Beta function is defined on the interval from 0 to 1. It is used when the integrand contains a power of t and a power of 1-t. The most important step is to compare the exponents in the integral with the exponents in the standard formula.",
    )
    formula_table(
        document,
        [
            ("Form", "Formula"),
            ("First form", "B(m,n) = ∫₀¹ t^(m-1) (1-t)^(n-1) dt"),
            ("Matching rule", "m - 1 = power of t,     n - 1 = power of (1-t)"),
        ],
    )
    para(
        document,
        "For example, in the integral ∫₀¹ t^0.5(1-t)^2 dt, we compare m - 1 = 0.5 and n - 1 = 2. Therefore m = 1.5 and n = 3, so the integral is written directly as B(1.5,3).",
    )
    para(
        document,
        "The second form is used for improper integrals from 0 to infinity. It is especially useful when the expression contains a factor like 1+t raised to a negative power. This form helps convert an infinite-limit integral into a Beta function value.",
    )
    formula_table(
        document,
        [
            ("Form", "Formula"),
            ("Second form", "B(m,n) = ∫₀∞ t^(m-1) (1+t)^-(m+n) dt"),
            ("Matching rule", "m - 1 = power of t,     m + n = exponent in (1+t)"),
        ],
    )
    para(
        document,
        "The third form is the trigonometric form. It is used when the integral includes powers of sine and cosine over the interval from 0 to π/2. This form is common in examples because the powers can be matched directly with 2m-1 and 2n-1.",
    )
    formula_table(
        document,
        [
            ("Form", "Formula"),
            ("Third form", "B(m,n) = 2∫₀^(π/2) cos^(2m-1)(t) sin^(2n-1)(t) dt"),
            ("Useful result", "∫₀^(π/2) cos^(2m-1)(t) sin^(2n-1)(t) dt = 1/2 B(m,n)"),
        ],
    )
    para(
        document,
        "These three forms may seem different, but they all describe the same special function. The difference is only in the shape of the integral. The first form is best for polynomial-like expressions on a finite interval. The second form is best for improper integrals. The third form is best for trigonometric powers. A correct solution begins by recognizing which family the given integral belongs to.",
    )
    para(
        document,
        "In lecture problems, the values of m and n are usually obtained by direct comparison. For example, if the power is written as m-1, then m is one more than that power. If the trigonometric power is written as 2m-1, then m is obtained by adding 1 and dividing by 2. This simple comparison is the key step in almost every Beta function example.",
    )
    page_break(document)

    heading(document, "The Problem and Method")
    para(
        document,
        "The main problem in using the Beta function is not memorizing the formula only. The real skill is deciding which form is suitable for a given integral. A student must first inspect the limits of integration, then study the shape of the integrand, and finally compare the powers with the standard Beta formula.",
    )
    heading(document, "A. The Problem", level=2)
    para(
        document,
        "Some integrals do not appear in standard Beta form immediately. The variable may be x instead of t, the integral may contain x², or the limits may require a change of variable. For this reason, substitution is often needed before applying the Beta function.",
    )
    bullet(document, "If the limits are from 0 to 1 and the integrand contains t and 1-t, the first form is usually suitable.")
    bullet(document, "If the limits are from 0 to infinity and the integrand contains 1+t or 1+x², the second form should be considered.")
    bullet(document, "If the integral contains sine and cosine powers on a π/2 interval, the trigonometric form is usually the correct choice.")
    bullet(document, "If half-integer values appear, the Gamma relation is useful because Γ(1/2)=√π.")
    heading(document, "B. The Method", level=2)
    para(
        document,
        "The method can be summarized in four steps. First, choose the suitable Beta form according to the limits and the integrand. Second, perform substitution if the integral is not already in standard form. Third, compare exponents to find m and n. Fourth, use the relation between Beta and Gamma functions when a numerical value is required.",
    )
    para(
        document,
        "For example, in ∫₀¹ (1-x²)^0.5 dx, we use t = x². Then dx is replaced by dt/(2t^0.5), and the integral becomes 1/2 ∫₀¹ t^(-1/2)(1-t)^0.5 dt. By comparison with the first form, m = 0.5 and n = 1.5, so the result is 1/2 B(0.5,1.5).",
    )
    para(
        document,
        "A useful way to avoid mistakes is to write the standard formula beside the given integral before starting the comparison. This makes it easier to see which factor corresponds to t^(m-1), which factor corresponds to (1-t)^(n-1), and whether an outside constant such as 1/2 appears after substitution. Many wrong answers happen because the substitution is correct but the outside multiplier is forgotten.",
    )
    para(
        document,
        "Another important point is the change of limits. When a new variable is used, the old limits must be converted to the new variable. In the example above, x=0 gives t=0 and x=1 gives t=1. Because the limits stay the same, the final expression becomes simple. In other examples, the limits may change from 0 to π/4 into 0 to π/2, which is exactly what is needed for the trigonometric Beta form.",
    )
    page_break(document)

    heading(document, "Relation Between Beta and Gamma Functions")
    para(
        document,
        "The relation between the Beta and Gamma functions is the most important computational result in this topic. It changes the value of B(m,n) into a quotient involving Gamma functions. This is especially helpful when m and n are positive integers or half-integers.",
    )
    formula_table(
        document,
        [
            ("Property", "Formula"),
            ("Main relation", "B(m,n) = Γ(m)Γ(n) / Γ(m+n)"),
            ("Symmetry", "B(m,n) = B(n,m)"),
            ("Useful Gamma values", "Γ(n) = (n-1)! for positive integers,     Γ(1/2) = √π"),
        ],
    )
    para(
        document,
        "A direct example from the lecture is B(4,3). Using the relation, B(4,3)=Γ(4)Γ(3)/Γ(7). Since Γ(4)=3!, Γ(3)=2!, and Γ(7)=6!, the final value is 3!×2!/6! = 1/60. This result is much faster than evaluating the original integral by expansion.",
    )
    para(
        document,
        "Another important example is B(0.5,1.5). We have B(0.5,1.5)=Γ(0.5)Γ(1.5)/Γ(2). Since Γ(0.5)=√π, Γ(1.5)=0.5√π, and Γ(2)=1!, the value is 0.5π. Therefore, 1/2 B(0.5,1.5)=π/4.",
    )
    para(
        document,
        "The symmetry property means that changing the order of the variables does not change the result. In other words, B(m,n)=B(n,m). This is useful when one order is easier to compare with the integral or easier to calculate.",
    )
    para(
        document,
        "The Gamma function also gives a simple rule for moving between values. The rule Γ(x+1)=xΓ(x) allows half-integer values to be calculated step by step. For example, Γ(3/2)=1/2 Γ(1/2)=1/2√π, and Γ(5/2)=3/2 Γ(3/2)=3/4√π. These identities explain why many Beta function answers contain π.",
    )
    para(
        document,
        "This relation is one reason why the Beta function is powerful. It connects an integral expression with factorial-like values. When m and n are integers, the answer often becomes a fraction involving factorials. When one or both variables are half-integers, the answer usually contains √π or π. Therefore, the Beta-Gamma relation turns integration into algebra.",
    )
    para(
        document,
        "There are also convergence conditions behind the formulas. In the first form, the integral is well behaved when the powers near 0 and 1 do not cause divergence. In the second form, the behavior near infinity must also be considered. In lecture-level problems, the given examples are usually selected so that the required Beta value is meaningful, but recognizing these conditions helps explain why the formulas are not applied blindly to every integral.",
    )
    para(
        document,
        "The Gamma relation is also useful for checking answers. If B(4,3)=1/60, the value is positive and small, which is reasonable because the integral contains powers that reduce the area under the curve. If B(0.5,0.5)=π, the answer is larger because both endpoints have integrable singular behavior. These checks make the final result more reliable.",
    )
    page_break(document)

    heading(document, "Numerical Simulation")
    para(
        document,
        "In this section, numerical simulation means following the computational steps symbolically and numerically: transform the integral, match it with a Beta form, and calculate the value using the Gamma relation. The following examples are based on the lecture style.",
    )
    heading(document, "Example 1: Integral on [0,1]", level=2)
    para(document, "I = ∫₀¹ (1 - x²)^0.5 dx", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(
        document,
        "Let t = x². The limits remain from 0 to 1, and dx becomes dt/(2t^0.5). Therefore, I = 1/2 ∫₀¹ t^(-1/2)(1-t)^0.5 dt. Comparing with the first Beta form gives m = 0.5 and n = 1.5. Thus, I = 1/2 B(0.5,1.5). Since B(0.5,1.5)=0.5π, the final result is I = π/4.",
    )
    heading(document, "Example 2: Trigonometric Integral", level=2)
    para(document, "I = ∫₀^(π/2) cos²(t) sin⁴(t) dt", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(
        document,
        "We compare the integrand with cos^(2m-1)(t) sin^(2n-1)(t). From cos²(t), 2m-1=2, so m=1.5. From sin⁴(t), 2n-1=4, so n=2.5. Since the trigonometric integral equals 1/2 B(m,n), we get I = 1/2 B(1.5,2.5).",
    )
    para(
        document,
        "Using the Gamma relation, B(1.5,2.5)=Γ(1.5)Γ(2.5)/Γ(4). With Γ(1.5)=0.5√π, Γ(2.5)=1.5×0.5√π, and Γ(4)=3!, the final value becomes π/32.",
    )
    heading(document, "Example 3: Lecture-Style Matching", level=2)
    para(document, "I = ∫₀¹ t^0.5(1-t)^2 dt", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(
        document,
        "This example is already in the first standard form. We compare t^0.5 with t^(m-1), so m-1=0.5 and m=1.5. We also compare (1-t)^2 with (1-t)^(n-1), so n-1=2 and n=3. Therefore, the integral is exactly B(1.5,3). This example shows the easiest case, where no substitution is needed before applying the Beta function.",
    )
    heading(document, "Example 4: Checking the Final Result", level=2)
    para(
        document,
        "After obtaining a Beta form, the final answer should be checked. The answer must be positive if the integrand is positive on the interval. The answer should also have a reasonable size when compared with the length of the interval and the maximum value of the integrand. For example, in ∫₀¹ (1-x²)^0.5 dx, the integrand is between 0 and 1, so the result must be between 0 and 1. The value π/4 is about 0.785, which is reasonable.",
    )
    para(
        document,
        "This checking step is important because Beta function problems contain many small details: exponents, limits, substitution constants, and the factor 1/2 in the trigonometric form. A correct method should produce not only a symbolic answer, but also an answer that makes sense numerically.",
    )
    page_break(document)

    heading(document, "Results and Applications")
    para(
        document,
        "The lecture examples show that the Beta function is not only a theoretical definition. It is a practical tool for solving definite integrals. It groups many different-looking integrals into three families: integrals on [0,1], improper integrals on [0,∞), and trigonometric integrals.",
    )
    heading(document, "Example 3: Improper Integral", level=2)
    para(document, "I = ∫₀∞ 1/(1+x²) dx", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    para(
        document,
        "Let t = x². Then the integral becomes 1/2 ∫₀∞ t^(-0.5)(1+t)^(-1) dt. Comparing with the second form gives m - 1 = -0.5, so m = 0.5. Also, m+n = 1, so n = 0.5. Therefore, I = 1/2 B(0.5,0.5). Since B(0.5,0.5)=π, the result is I = π/2.",
    )
    heading(document, "Important Notes", level=2)
    bullet(document, "The Beta function should be used only after the integral is matched with a standard form.")
    bullet(document, "Negative exponents are acceptable when they appear in the proper Beta form.")
    bullet(document, "For trigonometric integrals, remember that the integral is one half of B(m,n), not B(m,n) itself.")
    bullet(document, "The Gamma relation is the fastest way to calculate final values involving integers and half-integers.")
    para(
        document,
        "Applications of the Beta function appear in probability, statistics, physics, and engineering mathematics. In this course, its most direct use is evaluating definite integrals in a clean and organized way.",
    )
    para(
        document,
        "The main result from these applications is that different integrals can have the same hidden structure. A polynomial integral, an improper rational integral, and a trigonometric integral may look unrelated, but the Beta function provides one language for all of them. This is why special functions are useful in advanced mathematics: they collect repeated patterns into a single theory.",
    )
    para(
        document,
        "For a student, the practical benefit is speed and accuracy. Once the correct form is identified, the solution becomes shorter and less dependent on long algebraic manipulation. The student can also check the reasonableness of the answer by looking at the sign, the limits, and the expected size of the integral.",
    )
    para(
        document,
        "These results also show why the Beta function belongs to the study of complex functions and special functions. It does not solve only one isolated exercise; it gives a general framework. When a new integral appears, the student can ask whether it belongs to one of the known Beta families. If it does, the solution becomes direct and organized.",
    )
    page_break(document)

    heading(document, "CONCLUSION")
    para(
        document,
        "This paper explained the Beta function B(m,n) as a central special function in BAS115 Complex Functions. The Beta function has three useful forms: the first form on [0,1], the second form on [0,∞), and the trigonometric form on [0,π/2]. Each form is selected according to the limits and structure of the integral.",
    )
    para(
        document,
        "The most important conclusion is that solving Beta function problems depends on matching and substitution. After the integral is written in a standard form, the values of m and n can be identified directly. Then the relation B(m,n)=Γ(m)Γ(n)/Γ(m+n) gives the final numerical value in many cases.",
    )
    para(
        document,
        "Therefore, the Beta function is a powerful method because it reduces long integration work into a clear sequence of steps. This makes it useful for students studying complex functions and special functions.",
    )
    para(
        document,
        "The examples in this paper also show that the lecture method is consistent. Each solved problem follows the same pattern: choose the correct form, transform the variable if needed, compare exponents, and then use Gamma values. This repeated structure makes the topic easier to revise before exams because the same strategy can be applied to many questions.",
    )
    para(
        document,
        "In future study, the Beta function can be connected to probability distributions, Fourier-type integrals, and more advanced topics in mathematical physics. However, for BAS115, the most important achievement is understanding how this function simplifies definite integrals and why it is connected to the Gamma function.",
    )
    para(
        document,
        "Overall, the Beta function is valuable because it turns recognition into calculation. Once the form is recognized, the rest of the work follows from clear rules. This is the main lesson taken from Lecture 6.",
    )
    heading(document, "Acknowledgment")
    para(
        document,
        "The author thanks the BAS115 instructor for Lecture 6, which provided the main formulas, examples, and solution methods used in this paper.",
    )
    heading(document, "References")
    refs = [
        "M. Semary, BAS115 Complex Functions, Lecture 6: Beta Function, Spring 2025-2026.",
        "E. Kreyszig, Advanced Engineering Mathematics, 10th ed., Wiley, sections on Gamma and Beta functions.",
        "G. B. Thomas and M. D. Weir, Thomas' Calculus, Pearson, material on improper integrals and special functions.",
    ]
    for index, ref in enumerate(refs, 1):
        para(document, f"[{index}] {ref}", size=11.5, align=WD_ALIGN_PARAGRAPH.LEFT, after=3)
    heading(document, "Authors")
    para(
        document,
        "Student Name - BAS115 Complex Functions student, Department, Institute Name. Email: student@example.com",
        size=11.5,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        after=3,
    )
    para(
        document,
        "Correspondence Author - Student Name, student@example.com, contact number.",
        size=11.5,
        align=WD_ALIGN_PARAGRAPH.LEFT,
        after=3,
    )

    document.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
