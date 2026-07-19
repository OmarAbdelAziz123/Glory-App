from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


TEMPLATE = Path("/Users/omarabdelaziz/Downloads/CF-paper-template (3).docx")
OUT = Path("/Volumes/Extreme Pro/done_shop_projejcts/apps/glory_gym/Beta_Function_BAS115_Research.docx")


def clear_document(document: Document) -> None:
    body = document._body._element
    sect_pr = None
    for child in list(body):
        if child.tag == qn("w:sectPr"):
            sect_pr = child
        body.remove(child)
    if sect_pr is not None:
        body.append(sect_pr)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="B7B7B7", size="8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_rtl(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    bidi = p_pr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        p_pr.append(bidi)
    bidi.set(qn("w:val"), "1")


def set_run_font(run, font="Arial", size=10, bold=False, italic=False, color=None) -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)
    run._element.rPr.rFonts.set(qn("w:cs"), font)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def paragraph(
    document,
    text="",
    *,
    style=None,
    align=WD_ALIGN_PARAGRAPH.RIGHT,
    size=10.5,
    bold=False,
    italic=False,
    space_after=4,
    rtl=True,
    color=None,
):
    p = document.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.08
    if rtl:
        set_rtl(p)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color)
    return p


def ltr_paragraph(document, text="", *, size=10, bold=False, italic=False, space_after=4, align=WD_ALIGN_PARAGRAPH.LEFT):
    return paragraph(
        document,
        text,
        align=align,
        size=size,
        bold=bold,
        italic=italic,
        space_after=space_after,
        rtl=False,
    )


def heading(document, text, level=1):
    if level == 1:
        p = paragraph(document, text, size=12, bold=True, space_after=6, color="000000")
    else:
        p = paragraph(document, text, size=11, bold=True, space_after=4, color="000000")
    return p


def bullet(document, text):
    p = paragraph(document, "• " + text, size=10.2, space_after=3)
    p.paragraph_format.left_indent = Inches(0.16)
    return p


def page_break(document):
    p = document.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def add_formula_box(document, rows):
    table = document.add_table(rows=len(rows), cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.3)
    table.columns[1].width = Inches(4.8)
    for i, (name, formula) in enumerate(rows):
        c0, c1 = table.rows[i].cells
        c0.width = Inches(2.3)
        c1.width = Inches(4.8)
        set_cell_border(c0)
        set_cell_border(c1)
        if i == 0:
            set_cell_shading(c0, "EFEFEF")
            set_cell_shading(c1, "EFEFEF")
        for cell, text, align in ((c0, name, WD_ALIGN_PARAGRAPH.RIGHT), (c1, formula, WD_ALIGN_PARAGRAPH.LEFT)):
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = align
            if align == WD_ALIGN_PARAGRAPH.RIGHT:
                set_rtl(p)
            r = p.add_run(text)
            set_run_font(r, size=9.5, bold=(i == 0))
    document.add_paragraph()


def setup_document(document: Document) -> None:
    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.header_distance = Inches(0.35)
        section.footer_distance = Inches(0.35)
        header = section.header
        header.is_linked_to_previous = False
        header_p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        header_p.clear()
        header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = header_p.add_run("BAS115 Complex Functions Student Activity, Volume 2025-2026, Issue Spring, May 2026")
        set_run_font(r, size=8, italic=True, color="555555")
        footer = section.footer
        footer_p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        footer_p.clear()
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = footer_p.add_run("Beta Function Research")
        set_run_font(r, size=8, color="555555")

    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:cs"), "Arial")


def build():
    document = Document(str(TEMPLATE))
    clear_document(document)
    setup_document(document)

    paragraph(
        document,
        "دالة بيتا Beta Function في مادة الدوال المركبة",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=16,
        bold=True,
        space_after=8,
    )
    paragraph(
        document,
        "بحث مقدم إلى BAS115 Complex Functions Student Activity",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=11,
        bold=True,
        space_after=6,
    )
    paragraph(
        document,
        "إعداد: اسم الطالب    |    القسم: Department, Institute Name",
        align=WD_ALIGN_PARAGRAPH.CENTER,
        size=9.5,
        space_after=14,
    )
    ltr_paragraph(document, "Abstract-", bold=True, italic=True, size=10, space_after=0)
    paragraph(
        document,
        "يتناول هذا البحث دالة بيتا B(m,n) بوصفها واحدة من أهم الدوال الخاصة المستخدمة في تبسيط وحساب التكاملات المحددة في الدوال المركبة. يبدأ البحث بتعريف الدالة في صورتها الأساسية على الفترة من 0 إلى 1، ثم يعرض صورتها على الفترة من 0 إلى ما لا نهاية، وصورتها المثلثية المرتبطة بتكاملات الجيب وجيب التمام. كما يوضح البحث العلاقة المهمة بين دالة بيتا ودالة جاما، وكيف تساعد هذه العلاقة على حساب قيم كثيرة دون الرجوع إلى التكامل المباشر. ويعتمد البحث على أمثلة المحاضرة السادسة في BAS115 لشرح طريقة اختيار m و n وتحويل التكامل إلى صورة بيتا مناسبة.",
        size=10,
    )
    ltr_paragraph(document, "Index Terms-", bold=True, italic=True, size=10, space_after=0)
    paragraph(document, "Beta Function، Gamma Function، Special Functions، Definite Integrals، Complex Functions.", size=10)
    heading(document, "Introduction")
    paragraph(
        document,
        "تظهر الدوال الخاصة في الرياضيات عندما نحتاج إلى أدوات منظمة لحساب تكاملات أو متسلسلات لا يمكن التعامل معها بسهولة بالطرق الابتدائية. ومن هذه الدوال دالة بيتا، وهي دالة تعتمد على متغيرين وتربط بين الجبر والتحليل والتكاملات المحددة. أهميتها في هذه المحاضرة أنها تختصر خطوات طويلة في حساب التكامل، فبدل أن نبحث عن دالة أصلية مباشرة نستطيع مطابقة التكامل مع صورة قياسية ثم استخدام خصائص دالة بيتا.",
    )
    paragraph(
        document,
        "في مادة الدوال المركبة BAS115 تُدرس دالة بيتا بجانب دالة جاما لأنها تمثلان لغة موحدة للتعامل مع التكاملات التي تحتوي على قوى، حدود على الصورة 1 - t أو 1 + t، أو تكاملات مثلثية. لذلك فالهدف من البحث هو عرض الفكرة بطريقة مرتبة: تعريف الدالة، صورها الثلاث، العلاقة مع جاما، ثم تطبيقات محلولة مشابهة للمحاضرة.",
    )
    page_break(document)

    heading(document, "تعريف دالة بيتا والصور الأساسية")
    paragraph(
        document,
        "الصورة الأولى لدالة بيتا تعرف على الفترة [0,1]، وتستخدم عندما يكون التكامل مكتوبًا بدلالة t وقوى للحدين t و 1 - t. في هذه الحالة تتم المقارنة بين أسس التكامل والصورة القياسية لتحديد m و n.",
    )
    add_formula_box(
        document,
        [
            ("الصورة", "القانون"),
            ("الصورة الأولى", "B(m,n) = ∫₀¹ t^(m-1) (1-t)^(n-1) dt"),
            ("شروط المطابقة", "m - 1 = power of t,     n - 1 = power of (1-t)"),
        ],
    )
    paragraph(
        document,
        "مثال من المحاضرة: إذا كان التكامل ∫₀¹ t^0.5(1-t)^2 dt فإننا نساوي  m - 1 = 0.5  و  n - 1 = 2، وبالتالي m = 1.5 و n = 3. لذلك يكتب التكامل مباشرة على صورة B(1.5, 3).",
    )
    paragraph(
        document,
        "الصورة الثانية تستخدم مع تكاملات من 0 إلى ∞، خاصة عندما يظهر الحد (1+t) في المقام أو بقوة سالبة. وهذه الصورة مفيدة جدًا في التكاملات غير المنتهية لأنها تحول الحد اللانهائي إلى قيمة لدالة خاصة.",
    )
    add_formula_box(
        document,
        [
            ("الصورة", "القانون"),
            ("الصورة الثانية", "B(m,n) = ∫₀∞ t^(m-1) (1+t)^-(m+n) dt"),
            ("شروط المطابقة", "m - 1 = power of t,     m + n = exponent in the denominator"),
        ],
    )
    paragraph(
        document,
        "أما الصورة الثالثة فهي الصورة المثلثية، وتستخدم عندما يحتوي التكامل على قوى للجيب وجيب التمام خلال الفترة من 0 إلى π/2. وهي مهمة لأن كثيرًا من التكاملات المثلثية تصبح مباشرة بمجرد تحديد m و n.",
    )
    add_formula_box(
        document,
        [
            ("الصورة", "القانون"),
            ("الصورة الثالثة", "B(m,n) = 2∫₀^(π/2) cos^(2m-1)(t) sin^(2n-1)(t) dt"),
            ("استنتاج مهم", "∫₀^(π/2) cos^(2m-1)(t) sin^(2n-1)(t) dt = 1/2 B(m,n)"),
        ],
    )
    page_break(document)

    heading(document, "The problem and method")
    paragraph(
        document,
        "المشكلة الأساسية في موضوع دالة بيتا ليست حفظ القانون فقط، ولكن معرفة متى نستخدم كل صورة وكيف نحول التكامل إليها. لذلك تعتمد طريقة الحل على قراءة حدود التكامل أولًا، ثم فحص شكل الدالة داخل التكامل، وبعد ذلك مطابقة الأسس مع الصيغة المناسبة.",
    )
    heading(document, "A. The problem", level=2)
    paragraph(
        document,
        "قد يظهر التكامل في صورة مختلفة عن القانون القياسي. مثلًا قد يحتوي على x² بدل t، أو حدود تكامل من 0 إلى π/4 بدل 0 إلى π/2، أو يحتوي على tan و sin و cos في صورة مختلطة. هنا لا يكفي أن نرى تشابهًا عامًا، بل يجب عمل substitution مناسب حتى تصبح الحدود والأسس مطابقة لصورة بيتا.",
    )
    bullet(document, "إذا كانت الحدود من 0 إلى 1 وكان الشكل يحتوي على t و 1-t، نبدأ بالصورة الأولى.")
    bullet(document, "إذا كانت الحدود من 0 إلى ∞ وكان الشكل يحتوي على 1+t أو 1+x²، نبحث عن substitution يحول التكامل للصورة الثانية.")
    bullet(document, "إذا كان التكامل مثلثيًا على فترة تنتهي عند π/2، نستخدم الصورة الثالثة.")
    bullet(document, "إذا ظهرت قيم نصفية مثل 0.5 أو 1.5، نستخدم علاقة بيتا مع جاما لتسهيل الحساب.")
    heading(document, "B. The method", level=2)
    paragraph(
        document,
        "خطوات الحل المنهجي تكون كالتالي: نحدد الصورة المناسبة من حدود التكامل، ثم نكتب القانون القياسي تحت التكامل، ثم نقارن الأسس لاستخراج m و n، ثم نستخدم العلاقة مع دالة جاما إذا كان المطلوب حساب قيمة عددية. وإذا كان التكامل يحتاج تغيير متغير، نكتب التحويل وحدود التكامل الجديدة بوضوح قبل استخدام دالة بيتا.",
    )
    paragraph(
        document,
        "في مثال ∫₀¹ (1-x²)^0.5 dx نستخدم التحويل t = x². عندها dt = 2x dx و x = t^(1/2)، فيتحول التكامل إلى 1/2 ∫₀¹ t^(-1/2)(1-t)^0.5 dt. بمطابقة الصورة الأولى نحصل على 1/2 B(0.5, 1.5).",
    )
    page_break(document)

    heading(document, "العلاقة بين دالة بيتا ودالة جاما")
    paragraph(
        document,
        "العلاقة بين بيتا وجاما هي أهم نتيجة عملية في هذا الدرس، لأنها تحول قيمة دالة بيتا إلى حاصل ضرب وقسمة في دالة جاما. وتصبح الحسابات أسهل خصوصًا عندما تكون m و n أعدادًا صحيحة أو نصفية.",
    )
    add_formula_box(
        document,
        [
            ("العلاقة الأساسية", "B(m,n) = Γ(m)Γ(n) / Γ(m+n)"),
            ("خاصية التماثل", "B(m,n) = B(n,m)"),
            ("قيم مفيدة", "Γ(n) = (n-1)! for positive integers,     Γ(1/2) = √π"),
        ],
    )
    paragraph(
        document,
        "مثال مباشر من المحاضرة: B(4,3) = Γ(4)Γ(3) / Γ(7). وبما أن Γ(4)=3! و Γ(3)=2! و Γ(7)=6! فإن الناتج يساوي 3!×2!/6! = 1/60. هذا المثال يوضح قوة العلاقة؛ لأن حساب التكامل الأصلي مباشرة سيكون أطول من استخدام جاما.",
    )
    paragraph(
        document,
        "مثال آخر: B(0.5,1.5) = Γ(0.5)Γ(1.5) / Γ(2). وبما أن Γ(0.5)=√π و Γ(1.5)=0.5√π و Γ(2)=1!، إذن B(0.5,1.5)=0.5π. ومن هنا فإن 1/2 B(0.5,1.5)=π/4.",
    )
    paragraph(
        document,
        "خاصية التماثل تعني أن تبديل المتغيرين لا يغير قيمة الدالة، أي B(m,n)=B(n,m). هذه الخاصية مفيدة عند ترتيب القيم أو عندما تكون إحدى القيم أسهل في الحساب من الأخرى. وتظهر أيضًا في تكاملات الجيب وجيب التمام لأن تبادل القوى بينهما يؤدي إلى نفس نوع العلاقة.",
    )
    page_break(document)

    heading(document, "Numerical Simulation")
    paragraph(
        document,
        "في هذا القسم نطبق القوانين على أمثلة قريبة من المحاضرة. لا يعتمد الحل على برنامج حاسوبي، بل على محاكاة خطوات الحل العددية والرمزية: تحويل، مطابقة، ثم حساب باستخدام علاقة جاما وبيتا.",
    )
    heading(document, "Example 1: تكامل على الفترة من 0 إلى 1", level=2)
    ltr_paragraph(document, "I = ∫₀¹ (1 - x²)^0.5 dx", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(
        document,
        "نضع t = x²، فتظل الحدود من 0 إلى 1، ويتحول dx إلى dt/(2t^0.5). لذلك يصبح التكامل I = 1/2 ∫₀¹ t^(-1/2)(1-t)^0.5 dt. بالمقارنة مع الصورة الأولى نجد m = 0.5 و n = 1.5، وبالتالي I = 1/2 B(0.5,1.5). وباستخدام علاقة جاما تكون B(0.5,1.5)=0.5π، إذن I = π/4.",
    )
    heading(document, "Example 2: تكامل مثلثي", level=2)
    ltr_paragraph(document, "I = ∫₀^(π/2) cos²(t) sin⁴(t) dt", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(
        document,
        "نقارن مع الصورة المثلثية: cos^(2m-1)(t) sin^(2n-1)(t). من cos²(t) نحصل على 2m - 1 = 2، إذن m = 1.5. ومن sin⁴(t) نحصل على 2n - 1 = 4، إذن n = 2.5. وبما أن التكامل المثلثي يساوي 1/2 B(m,n)، إذن I = 1/2 B(1.5,2.5).",
    )
    paragraph(
        document,
        "يمكن حساب القيمة باستخدام جاما: B(1.5,2.5)=Γ(1.5)Γ(2.5)/Γ(4). وباستخدام Γ(1.5)=0.5√π و Γ(2.5)=1.5×0.5√π و Γ(4)=3! نحصل على قيمة نهائية تساوي π/32.",
    )
    page_break(document)

    heading(document, "نتائج وتطبيقات دالة بيتا")
    paragraph(
        document,
        "تؤكد أمثلة المحاضرة أن دالة بيتا ليست مجرد تعريف نظري، بل أداة حساب مباشرة. فهي تختصر التكاملات التي تبدو مختلفة في ثلاث عائلات: تكاملات الفترة [0,1]، التكاملات غير المنتهية، والتكاملات المثلثية. وبمجرد تحديد العائلة الصحيحة يصبح الحل منظمًا.",
    )
    heading(document, "Example 3: تكامل غير منته", level=2)
    ltr_paragraph(document, "I = ∫₀∞ 1/(1+x²) dx", size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    paragraph(
        document,
        "نستخدم t = x². عندها يتحول التكامل إلى 1/2 ∫₀∞ t^(-0.5)(1+t)^(-1) dt. بالمقارنة مع الصورة الثانية نجد m - 1 = -0.5 وبالتالي m = 0.5، كما أن m+n = 1 ولذلك n = 0.5. إذن I = 1/2 B(0.5,0.5). وباستخدام علاقة جاما نحصل على B(0.5,0.5)=π، وبالتالي I = π/2.",
    )
    heading(document, "ملاحظات مهمة عند الحل", level=2)
    bullet(document, "لا نستخدم دالة بيتا إلا بعد أن تصبح حدود التكامل وشكل الدالة مطابقين للصورة القياسية.")
    bullet(document, "القوى السالبة لا تعني أن الحل خطأ؛ غالبًا تظهر في الصورة الثانية أو بعد تحويل المتغير.")
    bullet(document, "القيم النصفية في جاما تعتمد على Γ(1/2)=√π، ثم نستخدم العلاقة Γ(x+1)=xΓ(x).")
    bullet(document, "في التكاملات المثلثية يجب الانتباه إلى معامل 2 في قانون بيتا؛ التكامل نفسه يساوي نصف B(m,n).")
    paragraph(
        document,
        "عمليًا، يمكن استعمال دالة بيتا في مسائل الاحتمالات، الإحصاء، الفيزياء الرياضية، وتحليل التكاملات في الهندسة. لكن ضمن نطاق هذه المحاضرة يتركز استخدامها في حساب التكاملات المحددة بطريقة أسرع وأكثر ترتيبًا.",
    )
    page_break(document)

    heading(document, "CONCLUSION")
    paragraph(
        document,
        "خلص البحث إلى أن دالة بيتا B(m,n) تمثل أداة مركزية في موضوع الدوال الخاصة داخل مادة الدوال المركبة. فهي تربط بين عدة أنواع من التكاملات من خلال ثلاث صور أساسية: صورة الفترة من 0 إلى 1، صورة الفترة من 0 إلى ما لا نهاية، والصورة المثلثية. كما أن علاقتها بدالة جاما تجعل حساب القيم النهائية أسهل، خصوصًا عندما تكون المعاملات أعدادًا صحيحة أو نصفية.",
    )
    paragraph(
        document,
        "أهم ما يجب على الطالب اكتسابه من هذا الدرس هو مهارة المطابقة والتحويل. فالحل الصحيح يبدأ من تحديد الصورة المناسبة، ثم اختيار substitution إذا احتاج التكامل لذلك، ثم استخراج m و n، وأخيرًا استخدام علاقة بيتا وجاما للوصول إلى النتيجة. بهذه الطريقة تصبح المسائل التي تبدو مختلفة في الشكل جزءًا من إطار واحد واضح.",
    )
    heading(document, "Acknowledgment")
    paragraph(
        document,
        "أتقدم بالشكر إلى دكتور المادة على شرح Lecture 6 في BAS115 Complex Functions، والتي كانت المصدر الأساسي لفكرة هذا البحث وأمثلته.",
        size=10,
    )
    heading(document, "References")
    refs = [
        "M. Semary, BAS115 Complex Functions, Lecture 6: Beta Function, Spring 2025-2026.",
        "E. Kreyszig, Advanced Engineering Mathematics, 10th ed., Wiley, sections on Gamma and Beta functions.",
        "G. B. Thomas and M. D. Weir, Thomas' Calculus, Pearson, material on improper integrals and special functions.",
    ]
    for i, ref in enumerate(refs, 1):
        ltr_paragraph(document, f"[{i}] {ref}", size=9.5, space_after=2)
    heading(document, "Authors")
    paragraph(
        document,
        "اسم الطالب – طالب مقرر BAS115 Complex Functions، Department, Institute Name. البريد الإلكتروني: student@example.com",
        size=9.5,
    )
    paragraph(
        document,
        "Correspondence Author – اسم الطالب، student@example.com، contact number.",
        size=9.5,
    )

    document.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    build()
